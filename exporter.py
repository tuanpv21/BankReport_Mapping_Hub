# -*- coding: utf-8 -*-
"""
Export Engine for BankReport Mapping Hub.
Exports mappings to:
  - Excel Spec (.xlsx) matching exact template:
    STT | Table | Column | DataType | Tên nghiệp vụ (Mô tả) | Ghi chú điều kiện | Bảng nguồn | Cột nguồn | Công thức tính
  - Word Document (.docx)
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import database

def export_to_excel(system_code=None, report_code=None, output_path=None):
    if not output_path:
        prefix = f"{system_code}_{report_code}" if report_code else (system_code or "ALL")
        output_path = os.path.join(os.path.dirname(__file__), f"Export_Mapping_{prefix}.xlsx")

    conn = database.get_db()
    cur = conn.cursor()

    query = """
        SELECT m.target_table, m.field_code, m.data_type, m.field_name_vi,
               m.notes, m.source_table, m.source_column, m.transformation_rule
        FROM mappings m
        WHERE 1=1
    """
    params = []
    if system_code and system_code != "ALL":
        query += " AND m.system_code = ?"
        params.append(system_code)
    if report_code and report_code != "ALL":
        query += " AND m.report_code = ?"
        params.append(report_code)
    query += " ORDER BY m.system_code, m.report_code, m.id"

    cur.execute(query, params)
    rows = cur.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping Spec"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    mono_font = Font(name="Consolas", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 9 standard columns matching user template
    headers = [
        "STT", "Table", "Column", "DataType", "Tên nghiệp vụ (Mô tả)",
        "Ghi chú điều kiện", "Bảng nguồn", "Cột nguồn", "Công thức tính"
    ]

    # Title row
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"BẢNG MA TRẬN MAPPING CHỈ TIÊU BÁO CÁO ({system_code or 'TT35 & CIC'})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Header row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[3].height = 26

    # Data rows
    for r_idx, r_data in enumerate(rows, 1):
        row_num = r_idx + 3
        ws.row_dimensions[row_num].height = 20

        # STT
        c_stt = ws.cell(row=row_num, column=1, value=r_idx)
        c_stt.alignment = align_center
        c_stt.font = data_font
        c_stt.border = thin_border

        for c_idx, val in enumerate(r_data, 2):
            cell = ws.cell(row=row_num, column=c_idx, value=val or "")
            cell.border = thin_border
            if c_idx in (2, 3, 4, 7, 8):
                cell.alignment = align_center if c_idx == 4 else align_left
                cell.font = mono_font if c_idx in (2, 3, 4, 7, 8) else data_font
            else:
                cell.alignment = align_left
                cell.font = data_font

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.freeze_panes = "A4"
    wb.save(output_path)
    return output_path

def export_to_word(system_code=None, report_code=None, output_path=None):
    if not output_path:
        prefix = f"{system_code}_{report_code}" if report_code else (system_code or "ALL")
        output_path = os.path.join(os.path.dirname(__file__), f"Doc_Spec_{prefix}.docx")

    conn = database.get_db()
    cur = conn.cursor()

    query = """
        SELECT m.system_code, m.report_code, r.report_name, m.target_table, m.field_code, 
               m.data_type, m.field_name_vi, m.notes, m.source_table, m.source_column, 
               m.transformation_rule, m.regulatory_ref
        FROM mappings m
        LEFT JOIN reports r ON m.report_code = r.report_code AND m.system_code = r.system_code
        WHERE 1=1
    """
    params = []
    if system_code and system_code != "ALL":
        query += " AND m.system_code = ?"
        params.append(system_code)
    if report_code and report_code != "ALL":
        query += " AND m.report_code = ?"
        params.append(report_code)
    query += " ORDER BY m.system_code, m.report_code, m.id"

    cur.execute(query, params)
    rows = cur.fetchall()
    conn.close()

    doc = docx.Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("TÀI LIỆU ĐẶC TẢ ÁNH XẠ NGHIỆP VỤ SANG KỸ THUẬT")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run(f"Hệ thống: {system_code or 'TT35 & CIC'} | Báo cáo: {report_code or 'Toàn bộ danh mục'}")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True

    doc.add_paragraph()

    reports_dict = {}
    for r in rows:
        rpt = r[1]
        if rpt not in reports_dict:
            reports_dict[rpt] = {
                "name": r[2] or rpt,
                "system": r[0],
                "items": []
            }
        reports_dict[rpt]["items"].append(r)

    for rpt_code, info in reports_dict.items():
        h2 = doc.add_heading(level=2)
        r_h2 = h2.add_run(f"Báo cáo [{info['system']}] {rpt_code} - {info['name']}")
        r_h2.font.name = "Times New Roman"
        r_h2.font.bold = True
        r_h2.font.color.rgb = RGBColor(15, 23, 42)

        tbl = doc.add_table(rows=1, cols=7)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        headers = ["STT", "Table", "Column", "DataType", "Tên Nghiệp vụ (Mô tả)", "Nguồn (Bảng.Cột)", "Công thức tính"]
        widths = [Inches(0.4), Inches(1.1), Inches(1.0), Inches(0.9), Inches(1.7), Inches(1.4), Inches(1.5)]

        hdr_cells = tbl.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            hdr_cells[idx].width = widths[idx]
            run = hdr_cells[idx].paragraphs[0].runs[0]
            run.font.name = "Times New Roman"
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            shading = docx.oxml.parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(docx.oxml.ns.nsdecls('w')))
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shading)

        for stt, item in enumerate(info['items'], 1):
            row_cells = tbl.add_row().cells
            src_str = f"{item[8]}.{item[9]}" if item[8] and item[9] else (item[8] or item[9] or "-")
            row_values = [
                str(stt),
                item[3] or "-",
                item[4] or "-",
                item[5] or "-",
                item[6] or "-",
                src_str,
                item[10] or item[7] or "-"
            ]
            for idx, val in enumerate(row_values):
                row_cells[idx].text = val
                row_cells[idx].width = widths[idx]
                if row_cells[idx].paragraphs[0].runs:
                    r = row_cells[idx].paragraphs[0].runs[0]
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8.5)

        doc.add_paragraph()

    doc.save(output_path)
    return output_path

def export_matrix_template(report_code, output_path=None):
    if not output_path:
        output_path = os.path.join(os.path.dirname(__file__), f"Template_Mapping_{report_code}.xlsx")

    conn = database.get_db()
    cur = conn.cursor()

    rpt = cur.execute("SELECT report_code, system_code, report_name, target_table, procedure_name, cycle, unit FROM reports WHERE report_code = ?", (report_code,)).fetchone()
    if not rpt:
        conn.close()
        raise ValueError(f"Report {report_code} not found")

    rpt_code, sys_code, rpt_name, tgt_tbl, proc_name, cycle, unit = rpt

    sql = """
        SELECT m.id, m.row_id, m.col_id, m.field_code, m.field_name_vi, m.gl_account,
               m.calc_method, m.source_table, m.source_column, m.transformation_rule,
               m.formula_expr, m.notes, m.target_table, m.target_column
        FROM mappings m
        WHERE m.report_code = ?
        ORDER BY m.id
    """
    rows = cur.execute(sql, (report_code,)).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Mapping_Config"
    ws2 = wb.create_sheet("Huong_Dan_Nghiep_Vu")

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    mono_font = Font(name="Consolas", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Sheet 1: Mapping_Config
    ws1.merge_cells("A1:N1")
    title_cell = ws1["A1"]
    title_cell.value = f"TEMPLATE CẤU HÌNH MAPPING CHỈ TIÊU & CÔNG THỨC - [{rpt_code}] {rpt_name}"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:N2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"Hệ thống: {sys_code} | Bảng đích: {tgt_tbl} | Chu kỳ: {cycle or 'Định kỳ'} | Đơn vị: {unit or 'VND'} | Thủ tục: {proc_name or '-'}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    headers = [
        "STT", "Mã dòng (Row ID)", "Mã cột (Col ID)", "Tên chỉ tiêu (Hàng)",
        "Tên cột số liệu (Cột)", "Số hiệu TK / GL Code", "Phương pháp tính",
        "Bảng nguồn Core / ODS", "Cột nguồn / Số dư", "Công thức / Biểu thức tính",
        "Điều kiện lọc / Ghi chú", "Mã trường (Field Code)", "Cột đích", "Bảng đích"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws1.row_dimensions[3].height = 26

    for idx, r in enumerate(rows, 1):
        row_num = idx + 3
        ws1.row_dimensions[row_num].height = 20

        # Extract values
        m_id, row_id, col_id, f_code, f_name, gl_acc, calc_m, src_tbl, src_col, rule, formula, notes, tgt_t, tgt_c = r

        vals = [
            idx,
            row_id or "-",
            col_id or "-",
            f_name or f_code,
            tgt_c or "-",
            gl_acc or "",
            calc_m or "GL_CONFIG",
            src_tbl or "",
            src_col or "",
            formula or rule or "",
            notes or "",
            f_code,
            tgt_c or "",
            tgt_t or ""
        ]

        for c_idx, val in enumerate(vals, 1):
            cell = ws1.cell(row=row_num, column=c_idx, value=val)
            cell.border = thin_border
            if c_idx in [1, 2, 3, 6, 7, 12, 13, 14]:
                cell.alignment = align_center
                cell.font = mono_font if c_idx in [2, 3, 6, 12, 13, 14] else data_font
            else:
                cell.alignment = align_left
                cell.font = data_font

    # Auto column width for ws1
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2]: continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)
    ws1.freeze_panes = "A4"

    # 2. Sheet 2: Huong_Dan_Nghiep_Vu
    ws2.merge_cells("A1:E1")
    h_title = ws2["A1"]
    h_title.value = "HƯỚNG DẪN ĐIỀN THÔNG TIN MAPPING VÀ PHƯƠNG PHÁP TÍNH TOÁN"
    h_title.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    h_title.alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[1].height = 26

    guide_headers = ["Phương pháp tính", "Khi nào sử dụng?", "Các cột cần điền", "Ví dụ cấu hình mẫu", "Quy cách sinh mã / Procedure"]
    for c_idx, gh in enumerate(guide_headers, 1):
        cell = ws2.cell(row=3, column=c_idx, value=gh)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws2.row_dimensions[3].height = 24

    guides = [
        (
            "GL_CONFIG",
            "Dùng cho các báo cáo cân đối kế toán, tài chính tính theo số dư tài khoản GL (ví dụ: Tiền mặt 1011, Ngoại tệ 1031, Vốn điều lệ 60, Quỹ 611,...)",
            "- Số hiệu TK / GL Code\n- Bảng nguồn (mặc định: t1080_tb_gl_bal_quy_doi)\n- Cột nguồn (NO_DAU_KY, CO_DAU_KY, NO_PHAT_SINH, CO_PHAT_SINH, NO_CUOI_KY, CO_CUOI_KY)",
            "Số hiệu TK: 1011\nBảng nguồn: t1080_tb_gl_bal_quy_doi\nCột nguồn: NO_DAU_KY (cho Dư nợ đầu kỳ)\nCông thức: a.gl_code LIKE '1011%'",
            "Sinh câu lệnh MERGE INTO theo GL_CODE hoặc nạp vào bảng cấu hình CAL_STEP1"
        ),
        (
            "CORE_TABLE",
            "Dùng khi dữ liệu chỉ tiêu lấy trực tiếp từ các bảng giao dịch Core Banking/ODS chi tiết (Khách hàng, khế ước, tài sản đảm bảo, CIC,...)",
            "- Bảng nguồn Core/ODS (ví dụ: ODS_OD_ACCOUNT, ODS_LOAN_ACCOUNT_HIST)\n- Cột nguồn (ví dụ: outstanding_balance, customer_id)\n- Điều kiện lọc / Ghi chú (WHERE clause)",
            "Bảng nguồn: ODS_OD_ACCOUNT\nCột nguồn: outstanding_balance\nĐiều kiện lọc: status = 'A' AND currency = 'VND'",
            "Sinh câu lệnh INSERT INTO ... SELECT hoặc câu lệnh query trực tiếp trong Procedure"
        ),
        (
            "FORMULA",
            "Dùng cho các chỉ tiêu tổng hợp hoặc tính toán đại số giữa các dòng/cột trong báo cáo",
            "- Công thức / Biểu thức tính toán (phép tính +, -, *, /, SUM, IF,...)",
            "(7) = (3) + (5)\nVỐN CẤP 1 (A) = A1 - A2\nLợi nhuận = TK69 + Total 7 - Total 8\nIF(TK69 + Total 7 - Total 8 > 0, ...)",
            "Sinh câu lệnh UPDATE tính toán tổng hợp hoặc subquery theo các dòng đã tính"
        ),
        (
            "PROCEDURE",
            "Dùng khi logic xử lý quá phức tạp hoặc đã được đóng gói thành các hàm/thủ tục PL/SQL riêng trong Package",
            "- Tên hàm hoặc Procedure xử lý (ví dụ: pr_A02211, pr_G04224, pk_cic_calc.fn_CIC_KU)\n- Ghi chú logic chi tiết",
            "Thủ tục: pr_A02211(p_date, v_cur)\nLogic: Chạy vòng lặp for cursor merge dữ liệu từ bảng quy đổi",
            "Được ghi nhận để liên kết vào file script .prc / .pck tương ứng"
        )
    ]

    for g_idx, g_row in enumerate(guides, 1):
        r_num = g_idx + 3
        ws2.row_dimensions[r_num].height = 65
        for col_i, val in enumerate(g_row, 1):
            cell = ws2.cell(row=r_num, column=col_i, value=val)
            cell.border = thin_border
            cell.alignment = align_left
            cell.font = data_font
            if col_i == 1:
                cell.font = Font(name="Consolas", size=11, bold=True, color="1E3A8A")
                cell.alignment = align_center

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 38
    ws2.column_dimensions["C"].width = 35
    ws2.column_dimensions["D"].width = 40
    ws2.column_dimensions["E"].width = 35

    wb.save(output_path)
    return output_path


def import_matrix_template(file_path):
    """Reads uploaded filled template Excel and updates mapping_hub.db."""
    if not os.path.exists(file_path):
        return {"success": False, "error": "File không tồn tại"}

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["Mapping_Config"] if "Mapping_Config" in wb.sheetnames else wb.active

    conn = database.get_db()
    cur = conn.cursor()

    updated = 0
    # Headers are at row 3, data from row 4
    for r in range(4, ws.max_row + 1):
        row_id = str(ws.cell(row=r, column=2).value or "").strip()
        col_id = str(ws.cell(row=r, column=3).value or "").strip()
        f_name = str(ws.cell(row=r, column=4).value or "").strip()
        gl_acc = str(ws.cell(row=r, column=6).value or "").strip()
        calc_m = str(ws.cell(row=r, column=7).value or "GL_CONFIG").strip().upper()
        src_tbl = str(ws.cell(row=r, column=8).value or "").strip()
        src_col = str(ws.cell(row=r, column=9).value or "").strip()
        formula = str(ws.cell(row=r, column=10).value or "").strip()
        notes = str(ws.cell(row=r, column=11).value or "").strip()
        f_code = str(ws.cell(row=r, column=12).value or "").strip()

        if not f_code and not (row_id and col_id):
            continue

        # Update matching mapping
        if f_code:
            cur.execute("""
                UPDATE mappings SET
                    gl_account = coalesce(nullif(?, ''), gl_account),
                    calc_method = coalesce(nullif(?, ''), calc_method),
                    source_table = coalesce(nullif(?, ''), source_table),
                    source_column = coalesce(nullif(?, ''), source_column),
                    transformation_rule = coalesce(nullif(?, ''), transformation_rule),
                    formula_expr = coalesce(nullif(?, ''), formula_expr),
                    notes = coalesce(nullif(?, ''), notes),
                    updated_at = CURRENT_TIMESTAMP
                WHERE field_code = ?
            """, (gl_acc, calc_m, src_tbl, src_col, formula, formula, notes, f_code))
            if cur.rowcount > 0:
                updated += 1
        elif row_id and col_id:
            cur.execute("""
                UPDATE mappings SET
                    gl_account = coalesce(nullif(?, ''), gl_account),
                    calc_method = coalesce(nullif(?, ''), calc_method),
                    source_table = coalesce(nullif(?, ''), source_table),
                    source_column = coalesce(nullif(?, ''), source_column),
                    transformation_rule = coalesce(nullif(?, ''), transformation_rule),
                    formula_expr = coalesce(nullif(?, ''), formula_expr),
                    notes = coalesce(nullif(?, ''), notes),
                    updated_at = CURRENT_TIMESTAMP
                WHERE row_id = ? AND col_id = ?
            """, (gl_acc, calc_m, src_tbl, src_col, formula, formula, notes, row_id, col_id))
            if cur.rowcount > 0:
                updated += 1

    conn.commit()
    conn.close()

    return {"success": True, "updated_count": updated}


def generate_sample_mapping_template(output_path=None):
    """Generates a clean, professional, standardized Excel template for users to download."""
    if not output_path:
        output_path = os.path.join(os.path.dirname(__file__), "Template_Mapping_Mau_Chuan.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping_Spec"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sample_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A1:I1")
    t_cell = ws["A1"]
    t_cell.value = "TEMPLATE MAPPING CHỈ TIÊU BÁO CÁO NGÂN HÀNG (DÙNG ĐỂ IMPORT VÀO HỆ THỐNG)"
    t_cell.font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    t_cell.alignment = align_center
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    sub_cell = ws["A2"]
    sub_cell.value = "Hướng dẫn: Điền thông tin theo các cột từ dòng 4 trở đi. Cột B (Table) và C (Column) là bắt buộc. Hệ thống tự động tạo báo cáo và chỉ tiêu mapping."
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    headers = [
        "STT", "Table (Bảng đích)", "Column (Mã chỉ tiêu)", "DataType (Kiểu DL)", 
        "Tên nghiệp vụ (Mô tả)", "Ghi chú / Điều kiện lọc", "Bảng nguồn (Source Table)", 
        "Cột nguồn (Source Column)", "Công thức tính / Transformation Rule"
    ]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = align_center
        c.border = border
    ws.row_dimensions[3].height = 28

    samples = [
        (1, "ODS_RPT_B01", "MA_CHI_NHANH", "VARCHAR2(20)", "Mã chi nhánh ngân hàng", "Chi nhánh cấp 1 & PGD", "ODS_BRANCH", "BRANCH_CODE", "BRANCH_CODE"),
        (2, "ODS_RPT_B01", "DU_NO_TIEN_MAT", "NUMBER(20,4)", "Dư nợ tài khoản tiền mặt", "Tài khoản 1011 quy đổi VND", "t1080_tb_gl_bal_quy_doi", "bal_amt", "SUM(bal_amt) WHERE gl_account = '1011' AND ccy = 'VND'"),
        (3, "ODS_RPT_B01", "LAI_PHAI_THU", "NUMBER(20,4)", "Số lãi dự thu tồn đọng", "Phát sinh lũy kế trong kỳ", "ODS_LN_INTEREST", "ACCRUED_INT", "NVL(ACCRUED_INT, 0)"),
        (4, "ODS_RPT_B01", "NHOM_NO", "VARCHAR2(10)", "Phân loại nhóm nợ CIC", "Theo quy định Thông tư 31/NHNN", "ODS_CIC_ACCOUNT", "CIC_GROUP", "CIC_GROUP")
    ]

    for r_idx, s in enumerate(samples, 4):
        ws.row_dimensions[r_idx].height = 22
        for c_idx, val in enumerate(s, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border
            cell.fill = sample_fill
            cell.alignment = align_center if c_idx in [1, 4] else align_left

    widths = [8, 22, 24, 18, 32, 28, 26, 22, 45]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    wb.save(output_path)
    return output_path


def import_new_report_excel(file_path, system_code=None, report_code=None, report_name=None, target_table=None, cycle="Tháng", overwrite=False):
    """Parses uploaded Excel file, creates or updates report in reports table, and inserts mappings."""
    if not os.path.exists(file_path):
        return {"success": False, "error": "File không tồn tại trên máy chủ"}

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["Mapping_Spec"] if "Mapping_Spec" in wb.sheetnames else (wb["Mapping_Config"] if "Mapping_Config" in wb.sheetnames else wb.active)

    header_row_idx = None
    col_map = {}
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, min(25, ws.max_column + 1))]
        filled = [v for v in row_vals if v]
        if len(filled) < 3:
            continue
        has_col = any("column" in v or "mã chỉ tiêu" in v or "mã cột" in v or "field" in v for v in row_vals)
        has_name = any("tên" in v or "mô tả" in v or "name" in v or "chỉ tiêu" in v for v in row_vals)
        if has_col or (has_name and any("bảng" in v or "table" in v for v in row_vals)):
            header_row_idx = r
            for c_idx, val in enumerate(row_vals, 1):
                if not val:
                    continue
                if "stt" in val or "no." in val:
                    col_map["stt"] = c_idx
                elif "table" in val or "bảng đích" in val:
                    col_map["target_table"] = c_idx
                elif "column" in val or "mã chỉ tiêu" in val or "mã cột" in val or "field_code" in val or "mã" in val:
                    if "target_column" not in col_map:
                        col_map["target_column"] = c_idx
                elif "datatype" in val or "kiểu dl" in val or "kiểu dữ liệu" in val or "data type" in val:
                    col_map["data_type"] = c_idx
                elif "tên" in val or "mô tả" in val or "diễn giải" in val or "field_name" in val:
                    col_map["field_name_vi"] = c_idx
                elif "ghi chú" in val or "điều kiện" in val or "note" in val:
                    col_map["notes"] = c_idx
                elif "bảng nguồn" in val or "source_table" in val or "source table" in val:
                    col_map["source_table"] = c_idx
                elif "cột nguồn" in val or "source_column" in val or "source column" in val:
                    col_map["source_column"] = c_idx
                elif "công thức" in val or "quy tắc" in val or "rule" in val or "formula" in val or "transformation" in val:
                    col_map["transformation_rule"] = c_idx
                elif "tài khoản" in val or "gl" in val or "account" in val:
                    col_map["gl_account"] = c_idx
                elif "phương pháp" in val or "method" in val:
                    col_map["calc_method"] = c_idx
            break

    if not header_row_idx:
        header_row_idx = 3
        col_map = {
            "stt": 1, "target_table": 2, "target_column": 3, "data_type": 4,
            "field_name_vi": 5, "notes": 6, "source_table": 7, "source_column": 8,
            "transformation_rule": 9
        }

    def get_val(r_num, key, default=""):
        c = col_map.get(key)
        if not c:
            return default
        v = ws.cell(row=r_num, column=c).value
        return str(v).strip() if v is not None else default

    data_rows = []
    found_tables = set()
    for r in range(header_row_idx + 1, ws.max_row + 1):
        f_code = get_val(r, "target_column")
        f_name = get_val(r, "field_name_vi")
        t_tbl = get_val(r, "target_table")

        if not f_code and not f_name:
            continue
        if not f_code and f_name:
            f_code = f"COL_{r - header_row_idx}"

        if t_tbl:
            found_tables.add(t_tbl.upper())

        data_rows.append({
            "target_table": t_tbl,
            "field_code": f_code,
            "field_name_vi": f_name or f_code,
            "data_type": get_val(r, "data_type", "VARCHAR2(50)"),
            "notes": get_val(r, "notes"),
            "source_table": get_val(r, "source_table"),
            "source_column": get_val(r, "source_column"),
            "transformation_rule": get_val(r, "transformation_rule"),
            "gl_account": get_val(r, "gl_account"),
            "calc_method": get_val(r, "calc_method")
        })

    if not data_rows:
        return {"success": False, "error": "Không tìm thấy dòng dữ liệu chỉ tiêu nào trong file Excel"}

    if not report_code:
        if found_tables:
            primary_tbl = list(found_tables)[0]
            clean_rpt = primary_tbl.replace("ODS_RPT_", "").replace("RPTB_", "").replace("CIC_", "").replace("RPT_", "")
            report_code = clean_rpt or primary_tbl
        else:
            base_fname = os.path.splitext(os.path.basename(file_path))[0]
            report_code = base_fname.replace("Template_", "").replace("Mapping_", "")

    report_code = str(report_code).strip().upper()

    if not system_code:
        if "CIC" in report_code or any("CIC" in t for t in found_tables):
            system_code = "CIC"
        elif "GL" in report_code or any("GL" in t for t in found_tables):
            system_code = "GL"
        else:
            system_code = "TT35"

    if not report_name:
        report_name = f"Báo cáo {report_code}"

    if not target_table:
        target_table = list(found_tables)[0] if found_tables else f"ODS_RPT_{report_code}"

    conn = database.get_db()
    cur = conn.cursor()

    existing_rpt = cur.execute("SELECT report_code FROM reports WHERE report_code = ? AND system_code = ?", (report_code, system_code)).fetchone()
    if not existing_rpt:
        cur.execute("""
            INSERT INTO reports (report_code, system_code, report_name, target_table, cycle, template_file, has_matrix)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (report_code, system_code, report_name, target_table, cycle or "Tháng", os.path.basename(file_path), 0))
    else:
        cur.execute("""
            UPDATE reports SET
                report_name = COALESCE(NULLIF(?, ''), report_name),
                target_table = COALESCE(NULLIF(?, ''), target_table),
                cycle = COALESCE(NULLIF(?, ''), cycle)
            WHERE report_code = ? AND system_code = ?
        """, (report_name, target_table, cycle or "Tháng", report_code, system_code))

    if overwrite:
        cur.execute("DELETE FROM mappings WHERE report_code = ? AND system_code = ?", (report_code, system_code))

    imported_count = 0
    for row in data_rows:
        row_target_tbl = row["target_table"] or target_table
        calc_method = row["calc_method"]
        if not calc_method:
            calc_method = "GL_CONFIG" if row["gl_account"] else "CORE_TABLE"

        cur.execute("""
            INSERT INTO mappings (
                system_code, report_code, field_code, field_name_vi, data_type,
                target_table, target_column, source_system, source_table, source_column,
                transformation_rule, notes, status, gl_account, calc_method
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'ODS', ?, ?, ?, ?, 'Active', ?, ?)
        """, (
            system_code,
            report_code,
            row["field_code"],
            row["field_name_vi"],
            row["data_type"],
            row_target_tbl,
            row["field_code"],
            row["source_table"],
            row["source_column"],
            row["transformation_rule"],
            row["notes"],
            row["gl_account"],
            calc_method
        ))
        imported_count += 1

    conn.commit()
    conn.close()

    return {
        "success": True,
        "system_code": system_code,
        "report_code": report_code,
        "report_name": report_name,
        "target_table": target_table,
        "imported_count": imported_count,
        "message": f"Nạp thành công {imported_count} chỉ tiêu cho báo cáo {report_code} ({system_code})"
    }
